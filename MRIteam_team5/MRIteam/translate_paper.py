import os
import time
import json
import re
import docx
from dotenv import load_dotenv
from google import genai

load_dotenv()
client = genai.Client()

original_file = "NCKH_Paper_Springer_VI.docx"
translated_file = "NCKH_Paper_Springer_EN.docx"

# Load the document
doc = docx.Document(original_file)

# Helper for calling Gemini with retry and backoff
def call_gemini_with_retry(client, model, contents, config=None, max_retries=10):
    for attempt in range(max_retries):
        try:
            # Add a small delay before every call to avoid burst requests
            time.sleep(2)
            response = client.models.generate_content(
                model=model,
                contents=contents,
                config=config
            )
            return response.text.strip()
        except Exception as e:
            print(f"Error calling Gemini (Attempt {attempt+1}/{max_retries}): {e}")
            if "429" in str(e) or "ResourceExhausted" in str(e) or "Quota" in str(e):
                # When rate limited, wait 70 seconds to completely clear the rolling 1-minute window
                wait_time = 70
                print(f"Rate limited. Waiting {wait_time}s to completely clear the window before retrying...")
                time.sleep(wait_time)
            else:
                time.sleep(5)
    raise Exception("Max retries exceeded for Gemini API call.")

# Parse HTML text containing <b> and <i> tags and apply to paragraph runs
def apply_styled_text_to_paragraph(p, html_text):
    html_text = html_text.replace('<strong>', '<b>').replace('</strong>', '</b>')
    html_text = html_text.replace('<em>', '<i>').replace('</em>', '</i>')
    
    # Save current alignment
    align = p.alignment
    
    # Clear existing runs
    p.text = ""
    
    tokens = re.split(r'(</?[bi]>)', html_text)
    is_bold = False
    is_italic = False
    
    for token in tokens:
        if token == '<b>':
            is_bold = True
        elif token == '</b>':
            is_bold = False
        elif token == '<i>':
            is_italic = True
        elif token == '</i>':
            is_italic = False
        elif token:
            run = p.add_run(token)
            run.bold = is_bold
            run.italic = is_italic
            
    p.alignment = align

# Individual paragraph translator
def translate_paragraph_individual(text):
    prompt = f"""You are a professional medical translator translating a scientific research paper on brain tumor classification from Vietnamese to English.
Translate clinical, anatomical, and machine learning terms accurately and professionally.

CRITICAL INSTRUCTIONS:
1. Preserve styling by surrounding bold text with <b>...</b> and italic text with <i>...</i>. If a word or phrase is bolded/italicized in Vietnamese, the corresponding English translation must be bolded/italicized using the same <b> or <i> tags.
2. Do NOT use markdown like ** or * or <strong> or <em>. Only use <b> and <i> tags.
3. If the text starts with "Hình X. ...", translate it as "Figure X. ...". If it starts with "Bảng Y. ...", translate it as "Table Y. ...".
4. Translate any references to figures (e.g. "(Hình X)" -> "(Figure X)") and tables (e.g. "(Bảng Y)" -> "(Table Y)") correctly.
5. Return ONLY the translated English HTML text. Do not include any explanation.

Text to translate:
{text}
"""
    return call_gemini_with_retry(client, 'gemini-2.5-flash-lite', prompt)

# Batch paragraph translator
def translate_paragraph_batch(batch):
    prompt = f"""You are a professional medical translator translating a scientific research paper on brain tumor classification from Vietnamese to English.
Here is a batch of paragraphs to translate represented as a JSON object:
{json.dumps(batch, ensure_ascii=False, indent=2)}

CRITICAL INSTRUCTIONS:
1. Translate all text values from Vietnamese to English. Keep clinical, anatomical, and deep learning terms accurate.
2. Preserve styling by surrounding bold text with <b>...</b> and italic text with <i>...</i>. If a word or phrase is bolded/italicized in Vietnamese, the corresponding English translation must be bolded/italicized using the same <b> or <i> tags.
3. Do NOT use markdown like ** or * or <strong> or <em>. Only use <b> and <i> tags.
4. If a paragraph starts with "Hình X. ...", translate it as "Figure X. ...". If it starts with "Bảng Y. ...", translate it as "Table Y. ...".
5. Translate all references to figures (e.g. "(Hình X)" -> "(Figure X)") and tables (e.g. "(Bảng Y)" -> "(Table Y)") correctly.
6. Keep the keys (e.g., "12", "15") exactly the same in the output JSON.
7. Return ONLY a valid JSON object. Do NOT wrap it in ```json ... ``` or include any explanation.
"""
    
    response_text = call_gemini_with_retry(client, 'gemini-2.5-flash-lite', prompt)
    
    if response_text.startswith("```json"):
        response_text = response_text[7:-3].strip()
    elif response_text.startswith("```"):
        response_text = response_text[3:-3].strip()
        
    try:
        translated_batch = json.loads(response_text)
        result = {}
        for k in batch.keys():
            if str(k) in translated_batch:
                result[int(k)] = translated_batch[str(k)]
            else:
                raise ValueError(f"Missing key {k}")
        return result
    except Exception as e:
        print(f"Batch translation failed: {e}. Falling back to individual translation...")
        result = {}
        for k, text in batch.items():
            time.sleep(3)
            result[k] = translate_paragraph_individual(text)
        return result

# Translate regular table cells in a batch
def translate_table_batch(table, table_idx):
    print(f"Translating Table {table_idx}...")
    table_data = {}
    for r_idx, row in enumerate(table.rows):
        table_data[f"row_{r_idx}"] = [cell.text for cell in row.cells]
        
    prompt = f"""You are translating a table from a medical machine learning paper from Vietnamese to English.
Here is the table represented as a JSON object:
{json.dumps(table_data, ensure_ascii=False, indent=2)}

CRITICAL INSTRUCTIONS:
1. Translate all text values from Vietnamese to English.
2. Ensure clinical terms are correct (e.g. 'Glioma', 'Meningioma', 'Pituitary', 'No Tumor', 'Accuracy', 'Precision', 'Recall', 'F1-Score', 'Support', 'Macro Average', 'Parameters', 'Overfitting', 'Severe', 'Very severe', 'Yes', 'No').
3. Convert all Vietnamese numbers to standard English formatting:
   - Convert decimal commas to decimal points (e.g., '96,11%' -> '96.11%', '1,08 s' -> '1.08 s').
   - Convert thousand separator dots to commas (e.g., '1.600' -> '1,600').
4. Keep the JSON keys (e.g. "row_0", "row_1") exactly the same.
5. Return ONLY a valid JSON object. Do NOT wrap it in ```json ... ``` or include any explanation.
"""
    
    response_text = call_gemini_with_retry(client, 'gemini-2.5-flash-lite', prompt)
    
    if response_text.startswith("```json"):
        response_text = response_text[7:-3].strip()
    elif response_text.startswith("```"):
        response_text = response_text[3:-3].strip()
        
    try:
        translated_data = json.loads(response_text)
        for r_idx, row in enumerate(table.rows):
            translated_row = translated_data[f"row_{r_idx}"]
            for c_idx, cell in enumerate(row.cells):
                bold = (r_idx == 0)
                align = None
                if len(cell.paragraphs) > 0:
                    align = cell.paragraphs[0].alignment
                cell.text = translated_row[c_idx]
                if len(cell.paragraphs) > 0:
                    p = cell.paragraphs[0]
                    if align is not None:
                        p.alignment = align
                    if bold and len(p.runs) > 0:
                        p.runs[0].font.bold = True
        print(f"Table {table_idx} translated successfully.")
    except Exception as e:
        print(f"Error parsing Table {table_idx} JSON: {e}. Translating cells individually...")
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                time.sleep(3)
                cell.text = translate_paragraph_individual(cell.text)
                if r_idx == 0 and len(cell.paragraphs) > 0 and len(cell.paragraphs[0].runs) > 0:
                    cell.paragraphs[0].runs[0].font.bold = True

# Translate code-like table cell comments
def translate_code_comments(table, table_idx):
    print(f"Translating code/comments in Table {table_idx}...")
    cell = table.cell(0, 0)
    original_code = cell.text
    
    prompt = f"""You are a software engineer translating comments inside Python code/pseudocode/JSON from Vietnamese to English.
Translate ONLY the comments (lines starting with # or containing //).
Do NOT translate any code keywords, variable names, function names, or programming syntax.
Return the complete code/pseudocode with only the comments translated.

Code/Pseudocode to process:
{original_code}
"""
    translated_code = call_gemini_with_retry(client, 'gemini-2.5-flash-lite', prompt)
    cell.text = translated_code
    print(f"Table {table_idx} code comments translated successfully.")

# ----------------------------------------------------------------------
# STEP 1: PARAGRAPH TRANSLATION
# ----------------------------------------------------------------------
print("Starting paragraph translation...")
paras_to_translate = {}
for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        paras_to_translate[idx] = p.text

# Group paragraphs into batches of 10
batch_size = 10
keys = list(paras_to_translate.keys())
batches = [keys[i:i + batch_size] for i in range(0, len(keys), batch_size)]

print(f"Total non-empty paragraphs: {len(keys)}")
print(f"Divided into {len(batches)} batches.")

translated_paras = {}
for batch_idx, batch_keys in enumerate(batches):
    print(f"Processing Batch {batch_idx+1}/{len(batches)}...")
    batch_data = {k: paras_to_translate[k] for k in batch_keys}
    
    translated_batch = translate_paragraph_batch(batch_data)
    translated_paras.update(translated_batch)
    
    # Sleep 6 seconds between batches to stay under rate limits
    time.sleep(6)

# Apply translated text to paragraphs
print("Applying translated text to document paragraphs...")
for idx, new_text in translated_paras.items():
    apply_styled_text_to_paragraph(doc.paragraphs[idx], new_text)

# ----------------------------------------------------------------------
# STEP 2: TABLE TRANSLATION
# ----------------------------------------------------------------------
print("\nStarting table translation...")
for idx, table in enumerate(doc.tables):
    table_idx = idx + 1
    if len(table.rows) == 1 and len(table.columns) == 1:
        translate_code_comments(table, table_idx)
    else:
        translate_table_batch(table, table_idx)
    # Sleep 6 seconds between tables to stay under rate limits
    time.sleep(6)

# ----------------------------------------------------------------------
# STEP 3: IMAGE REPLACEMENT
# ----------------------------------------------------------------------
print("\nReplacing figure images with English versions...")
image_mapping = {
    1: "models/preprocessing_pipeline_en.png",
    2: "models/brain_tumor_samples_en.png",
    4: "models/latency_analysis_en.png",
    6: "models/conflict_cases_en.png",
    7: "models/consensus_comparison_chart_en.png",
    8: "models/error_distribution_en.png",
    9: "models/confusion_matrix_consensus_en.png",
    10: "models/sota_comparison_en.png",
    11: "models/multimodal_explanation_en.png"
}

for idx, path in image_mapping.items():
    if os.path.exists(path):
        try:
            shape = doc.inline_shapes[idx]
            rId = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            if rId is not None:
                part = doc.part.related_parts[rId]
                with open(path, "rb") as f:
                    part._blob = f.read()
                print(f"Successfully replaced Image {idx} (rId={rId}) with {path}")
            else:
                print(f"Error replacing Image {idx}: rId is None")
        except Exception as e:
            print(f"Error replacing Image {idx}: {e}")
    else:
        print(f"Warning: English figure file not found at {path}")

# Save the document
print(f"\nSaving translated document to {translated_file}...")
doc.save(translated_file)
print("English translation compiled successfully!")
